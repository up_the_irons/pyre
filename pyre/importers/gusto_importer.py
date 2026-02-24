"""Gusto General Ledger XLSX parser."""

import re
from datetime import datetime
from pathlib import Path

import openpyxl
import yaml

from pyre.importers.journal import JournalEntry, JournalSplit, compute_journal_hash


_HEADER_ROW = ["Account Type", "Account Description", "Debit", "Credit"]


def _find_header_row(rows):
    """Return the index of the header row, or -1 if not found."""
    for i, row in enumerate(rows):
        if list(row) == _HEADER_ROW:
            return i
    return -1


def detect_gusto_gl(filepath):
    """Return True if the XLSX looks like a Gusto General Ledger export.

    Checks for a 'Basic' sheet containing a header row with:
    Account Type, Account Description, Debit, Credit.
    """
    try:
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    except Exception:
        return False

    try:
        if "Basic" not in wb.sheetnames:
            return False
        ws = wb["Basic"]
        rows = list(ws.iter_rows(values_only=True))
        return _find_header_row(rows) >= 0
    finally:
        wb.close()


def load_gusto_config(config_path=None):
    """Load the Gusto account mapping from gusto.yaml next to the database.

    Returns the parsed YAML dict (with an 'account_map' key).
    Raises FileNotFoundError if the file is missing.
    """
    if config_path is None:
        from pyre.db import DB_PATH
        config_path = DB_PATH.with_name("gusto.yaml")
    else:
        config_path = Path(config_path)

    if not config_path.exists():
        raise FileNotFoundError(
            f"Gusto config not found: {config_path}\n"
            "Copy gusto.yaml.sample next to your database "
            "and edit the account mappings."
        )

    with open(config_path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def parse_gusto_gl(filepath):
    """Parse a Gusto General Ledger XLSX into JournalEntry objects.

    Reads the 'Basic' sheet. Each file produces one JournalEntry
    (one pay period). The check date becomes the transaction date.

    Returns list of JournalEntry (always 0 or 1 element).
    """
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    try:
        ws = wb["Basic"]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    header_idx = _find_header_row(rows)
    if header_idx < 0:
        return []

    # Scan rows above the header for the description and check date.
    # These are the last two non-blank rows before the header.
    description = ""
    check_date_raw = ""
    for row in rows[:header_idx]:
        val = row[0]
        if val is None:
            continue
        val = str(val)
        if re.search(r"Check date:", val):
            check_date_raw = val
        elif re.search(r"Ledger for", val):
            description = val

    # Extract the ISO date from "Check date: YYYY-MM-DD"
    m = re.search(r"Check date:\s*(\d{4}-\d{2}-\d{2})", check_date_raw)
    if not m:
        return []
    check_date = m.group(1)

    # Validate the date parses
    try:
        datetime.strptime(check_date, "%Y-%m-%d")
    except ValueError:
        return []

    # Data starts after the header row
    splits = []
    for row in rows[header_idx + 1:]:
        account_type = row[0]
        if account_type is None:
            continue  # skip Totals row and blanks

        account_desc = row[1] or ""
        debit = row[2]
        credit = row[3]

        if debit is not None:
            amount_cents = int(round(float(debit) * 100))
        elif credit is not None:
            amount_cents = -int(round(float(credit) * 100))
        else:
            continue

        splits.append(JournalSplit(
            account_name=str(account_type),
            amount_cents=amount_cents,
            description=str(account_desc),
        ))

    if not splits:
        return []

    entry = JournalEntry(
        date=check_date,
        description=str(description),
        splits=splits,
    )
    entry.hash = compute_journal_hash(entry)
    return [entry]


def resolve_config_paths(account_map, path_map):
    """Convert account path values in the config to Pyre account IDs.

    account_map: raw dict from gusto.yaml (values are account paths or
                 dicts of description -> account path).
    path_map: dict from build_account_path_map() (path -> account ID).

    Returns (resolved_map, bad_paths) where resolved_map has the same
    structure but with account IDs instead of paths, and bad_paths is a
    set of paths that could not be resolved.
    """
    resolved = {}
    bad_paths = set()

    for gusto_type, mapping in account_map.items():
        if isinstance(mapping, str):
            acct_id = path_map.get(mapping)
            if acct_id:
                resolved[gusto_type] = acct_id
            else:
                bad_paths.add(mapping)
        elif isinstance(mapping, dict):
            sub = {}
            for desc_pattern, path in mapping.items():
                acct_id = path_map.get(path)
                if acct_id:
                    sub[desc_pattern] = acct_id
                else:
                    bad_paths.add(path)
            resolved[gusto_type] = sub

    return resolved, bad_paths


def _resolve_split(split, account_map):
    """Resolve a single split using the account map.

    The map value can be:
      - A string: all splits of this type map to that Pyre account ID.
      - A dict: keys are substrings matched against the split description.
                The first matching key wins.

    Returns the Pyre account ID, or None if unmatched.
    """
    mapping = account_map.get(split.account_name)
    if mapping is None:
        return None

    if isinstance(mapping, str):
        return mapping

    # Dict: match description substrings (case-insensitive)
    desc_lower = split.description.lower()
    for pattern, pyre_id in mapping.items():
        if pattern.lower() in desc_lower:
            return pyre_id

    return None


def resolve_gusto_accounts(entries, account_map):
    """Resolve Gusto account types to Pyre account IDs.

    account_map values can be a string (one account for the whole type) or
    a dict mapping description substrings to different accounts.

    Returns set of unmapped split labels (type or type:description).
    """
    unmatched = set()

    for entry in entries:
        for split in entry.splits:
            pyre_id = _resolve_split(split, account_map)
            if pyre_id:
                split.account_id = pyre_id
            else:
                label = split.account_name
                if split.description:
                    label = f"{label}: {split.description}"
                unmatched.add(label)

    return unmatched
